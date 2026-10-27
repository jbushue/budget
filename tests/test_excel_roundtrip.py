"""
Excel export -> import round-trip regression tests.

Builds a rich FinanceManager in memory (no database), exports it to a temp
xlsx, imports into a fresh FinanceManager and compares every field the Excel
format covers.  Payment state (paid-this-period flags, payment/transaction
history, carryover balances) is intentionally not part of the format.
"""

import os
from datetime import date

import pytest
from openpyxl import load_workbook

from models.bank import Bank
from models.bill_group import BillGroup, BillGroupMember
from models.data_import_export import DataImportExport, EXAMPLE_COLUMN
from models.finance_manager import FinanceManager
from models.fixed_bill import Bill, FloatingBill, GroupLinkedBill, AccountLinkedBill
from models.income_source import IncomeSource
from models.mortgage_account import EscrowBreakdown
from models.types import (
    AllocationMethod, BillFrequency, CreditAccountStatus, MortgageStatus,
)


def build_manager(first_bank_name: str = "Chase Bank") -> FinanceManager:
    b = Bank(first_bank_name, "traditional")
    b.add_account("checking", "Checking", 5000, 0.1, 5)
    b.add_account("savings", "Savings", 10000, 4.5, 1)

    cc = b.add_credit_account("Visa", 10000, 22.9, 20, 25)
    cc.balance = 3000
    cc.regular_balance = 1000
    cc.status = CreditAccountStatus.FROZEN
    p1 = cc.add_promotional_offer("BT 0%", 18, 5000, date(2026, 1, 1))
    p1.amount_used = 1500
    cc.promo_balance["BT 0%"] = 1500
    p2 = cc.add_promotional_offer("Purchase 0%", 12, None, date(2026, 3, 1))
    p2.amount_used = 600
    cc.promo_balance["Purchase 0%"] = 500
    p2.interest_charged = True

    loan = b.add_credit_account("Loan A", 20000, 5.0, 10, 1)
    loan.balance = 8000
    loan.regular_balance = 8000

    m = b.add_mortgage_account("Home", 400000, 6.5, 360, 350000, 1,
                               EscrowBreakdown(300, 120, 50, 25, 10),
                               first_bank_name, "Checking")
    m.start_date = date(2020, 5, 1)
    m.market_value = 500000
    m.escrow_balance = 1234.5
    m.status = MortgageStatus.REFINANCED

    b.add_heloc_account("HELOC", 50000, 8.0, 120, 240, date(2024, 2, 1), 12000, 15,
                        first_bank_name, "Checking")

    lender = Bank("Other Lender", "credit_issuer")
    loan_b = lender.add_credit_account("Loan B", 15000, 4.0, 10, 1)
    loan_b.balance = 5000
    loan_b.regular_balance = 5000

    bills = [
        Bill("Gym", 50, 3, frequency="monthly",
             preferred_source_bank=first_bank_name, preferred_source_account="Checking"),
        Bill("Property Tax", 1200, 15, frequency="yearly", due_month=4),
        Bill("Water", 90, 7, frequency="quarterly", start_date=date(2026, 2, 7)),
    ]
    el = FloatingBill("Electric", 150, 12, 80, 250, frequency="monthly")
    el.amount_history = [140.0, 160.0]
    el.set_actual_amount(175.0)
    bills.append(el)

    incomes = [
        IncomeSource("Salary", 3500, "semimonthly", [15, 31],
                     {"Checking": 0.9, "Savings": {"value": 0.5, "type": "fixed_amount"}}),
        IncomeSource("Bonus", 1000, "biweekly", [1], {"Savings": None},
                     start_date=date(2026, 1, 9)),
    ]

    groups = [BillGroup(
        "Student Loans", 300, 10, BillFrequency.MONTHLY, "student_loan",
        first_bank_name, "Checking",
        [BillGroupMember(first_bank_name, "Loan A", AllocationMethod.FIXED, 200.0),
         BillGroupMember("Other Lender", "Loan B", AllocationMethod.PROPORTIONAL, 0.0)],
    )]
    return FinanceManager([b, lender], incomes, bills, groups)


def snapshot(fm: FinanceManager) -> dict:
    out = {}
    for b in fm.banks:
        out[f"bank:{b.name}"] = {"bank_type": b.bank_type}
        for n, a in b.accounts.items():
            out[f"acct:{b.name}:{n}"] = dict(type=a.account_type, bal=a.balance, rate=a.interest_rate,
                                             accr=a.interest_accrual_date)
        for n, c in b.credit_accounts.items():
            out[f"credit:{b.name}:{n}"] = dict(
                limit=c.credit_limit, bal=c.balance, reg=c.regular_balance, rate=c.interest_rate,
                due=c.payment_due_date, stmt=c.statement_date, status=c.status,
                promo_balance=dict(c.promo_balance))
            for p in c.active_promos:
                out[f"promo:{n}:{p.promo_name}"] = dict(
                    months=p.months_interest_free, limit=p.purchase_limit, used=p.amount_used,
                    start=p.start_date, end=p.end_date, charged=p.interest_charged)
        for n, m in b.mortgage_accounts.items():
            out[f"mort:{b.name}:{n}"] = dict(
                orig=m.original_principal, cur=m.current_principal, rate=m.interest_rate,
                term=m.loan_term_months, due=m.payment_due_date, start=m.start_date,
                escrow=vars(m.escrow), escbal=m.escrow_balance,
                src=(m.preferred_source_bank, m.preferred_source_account),
                mv=m.market_value, status=m.status)
        for n, h in b.heloc_accounts.items():
            out[f"heloc:{b.name}:{n}"] = dict(
                limit=h.credit_limit, bal=h.balance, rate=h.interest_rate,
                draw=h.draw_period_months, term=h.total_term_months, start=h.start_date,
                due=h.payment_due_date, src=(h.preferred_source_bank, h.preferred_source_account))
    for x in fm.bills:
        d = dict(cls=type(x).__name__, amount=round(x.amount, 2), due=x.due_date,
                 due_month=x.due_month, type=x.bill_type, freq=x.frequency,
                 src=(x.preferred_source_bank, x.preferred_source_account),
                 start=getattr(x, "start_date", None))
        if isinstance(x, FloatingBill):
            d.update(est=x.estimated_amount, actual=x._actual_amount, mn=x.min_amount,
                     mx=x.max_amount, ahist=x.amount_history)
        out[f"bill:{x.bill_name}"] = d
    for i in fm.income_sources:
        out[f"income:{i.source_name}"] = dict(
            amount=i.amount, freq=i.frequency, pay=i.pay_date,
            start=getattr(i, "start_date", None),
            dist=i.distribution.to_json() if i.distribution else None)
    for g in fm.bill_groups:
        out[f"group:{g.group_name}"] = dict(
            total=g.total_payment, due=g.due_date, freq=g.frequency, type=g.group_type,
            src=(g.preferred_source_bank, g.preferred_source_account),
            members=[(m.bank_name, m.account_name, m.allocation_method, m.fixed_amount)
                     for m in g.members])
    return out


def diff(a: dict, b: dict) -> list[str]:
    problems = []
    for key in sorted(set(a) | set(b)):
        if key not in b:
            problems.append(f"missing after import: {key}")
        elif key not in a:
            problems.append(f"extra after import: {key}")
        else:
            for f in sorted(set(a[key]) | set(b[key])):
                if a[key].get(f) != b[key].get(f):
                    problems.append(f"{key}.{f}: {a[key].get(f)!r} != {b[key].get(f)!r}")
    return problems


def export_import(fm, tmp_path, target=None):
    path = os.path.join(tmp_path, "export.xlsx")
    DataImportExport(fm).export_current_data(path)
    target = target if target is not None else FinanceManager([], [], [], [])
    results = DataImportExport(target).import_from_excel(path)
    return target, results, path


@pytest.mark.parametrize("bank_name", ["Chase Bank", "Chase"])
def test_full_round_trip_has_no_diffs(tmp_path, bank_name):
    fm = build_manager(bank_name)
    before = snapshot(fm)
    fm2, results, _ = export_import(fm, tmp_path)
    assert results["success"], results["errors"]
    assert results["errors"] == []
    assert diff(before, snapshot(fm2)) == []


def test_round_trip_is_stable_twice(tmp_path):
    fm = build_manager()
    fm2, _, _ = export_import(fm, tmp_path)
    fm3, results, _ = export_import(fm2, os.path.join(tmp_path))
    assert results["success"]
    assert diff(snapshot(fm), snapshot(fm3)) == []


def test_group_linked_bill_not_exported_as_plain_bill(tmp_path):
    fm = build_manager()
    fm2, results, path = export_import(fm, tmp_path)
    ws = load_workbook(path)["Bills"]
    names = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]
    assert not any(str(n).startswith("[Auto]") for n in names if n)
    group_bills = [b for b in fm2.bills if b.bill_name == "[Auto] Student Loans Payment"]
    assert len(group_bills) == 1 and isinstance(group_bills[0], GroupLinkedBill)
    # Grouped credit accounts don't get their own linked bill
    assert not any(isinstance(b, AccountLinkedBill) and b.bill_name == "[Auto] Loan A Payment"
                   for b in fm2.bills)


def test_promo_dates_and_regular_balance(tmp_path):
    fm = build_manager()
    fm2, _, _ = export_import(fm, tmp_path)
    visa = fm2.banks[0].credit_accounts["Visa"]
    promos = {p.promo_name: p for p in visa.active_promos}
    assert promos["BT 0%"].start_date == date(2026, 1, 1)
    assert promos["Purchase 0%"].start_date == date(2026, 3, 1)
    assert visa.regular_balance == pytest.approx(3000 - 1500 - 500)


def test_reimport_into_same_manager_does_not_duplicate(tmp_path):
    fm = build_manager()
    before = snapshot(fm)
    n_bills, n_income = len(fm.bills), len(fm.income_sources)
    _, results, _ = export_import(fm, tmp_path, target=fm)
    assert results["errors"] == []
    assert len(fm.bills) == n_bills
    assert len(fm.income_sources) == n_income
    assert results["bills_created"] == 0
    assert results["bills_skipped"] == 4
    assert results["income_sources_skipped"] == 2
    assert results["bank_accounts_skipped"] == 2
    assert results["promos_skipped"] == 2
    assert results["bill_groups_skipped"] == 1
    assert diff(before, snapshot(fm)) == []


def test_template_example_rows_are_skipped(tmp_path):
    path = os.path.join(tmp_path, "template.xlsx")
    DataImportExport().generate_template(path)
    fm = FinanceManager([], [], [], [])
    results = DataImportExport(fm).import_from_excel(path)
    assert results["errors"] == []
    assert results["example_rows_skipped"] > 0
    assert fm.banks == [] and fm.income_sources == [] and fm.bill_groups == []
    assert [b for b in fm.bills] == []


def test_template_rows_import_once_example_flag_cleared(tmp_path):
    path = os.path.join(tmp_path, "template.xlsx")
    DataImportExport().generate_template(path)
    wb = load_workbook(path)
    for ws in wb.worksheets[1:]:
        headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
        col = headers[EXAMPLE_COLUMN]
        for r in range(3, ws.max_row + 1):
            ws.cell(row=r, column=col).value = None
    wb.save(path)

    fm = FinanceManager([], [], [], [])
    results = DataImportExport(fm).import_from_excel(path)
    assert results["errors"] == []
    assert {b.name for b in fm.banks} == {"Chase Bank", "Discover", "Wells Fargo Mortgage"}
    assert results["bank_accounts_created"] == 2
    assert results["helocs_created"] == 1
    assert results["bill_groups_created"] == 1
    assert results["promos_created"] == 1
    assert fm.banks[1].credit_accounts["Discover It"].regular_balance == pytest.approx(500.50)
    assert any(isinstance(b, GroupLinkedBill) for b in fm.bills)


def test_legacy_export_with_end_date_column(tmp_path):
    """Files written by the old exporter had only end_date in column 6."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Banks"
    ws.append(["bank_name", "bank_type"])
    ws.append(["(data exported)", ""])
    ws.append(["Chase Bank", "traditional"])
    ws = wb.create_sheet("Credit_Accounts")
    ws.append(["bank_name", "account_name", "credit_limit", "balance", "interest_rate",
               "payment_due_date", "statement_date"])
    ws.append(["(data exported)"])
    ws.append(["Chase Bank", "Visa", 10000, 2000, 20.0, 5, 1])
    ws = wb.create_sheet("Promo_Offers")
    ws.append(["credit_account_name", "promo_name", "months_interest_free",
               "purchase_limit", "current_balance", "end_date"])
    ws.append(["(data exported)"])
    ws.append(["Visa", "Old Promo", 12, None, 800, "2027-01-01"])
    path = os.path.join(tmp_path, "legacy.xlsx")
    wb.save(path)

    fm = FinanceManager([], [], [], [])
    results = DataImportExport(fm).import_from_excel(path)
    assert results["errors"] == []
    assert fm.banks[0].name == "Chase Bank"
    promo = fm.banks[0].credit_accounts["Visa"].active_promos[0]
    assert promo.end_date == date(2027, 1, 1)
    assert promo.start_date < promo.end_date
    assert fm.banks[0].credit_accounts["Visa"].regular_balance == pytest.approx(1200)
