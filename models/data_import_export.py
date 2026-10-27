"""
Data Import/Export module for bulk loading financial data from Excel.
Supports importing banks, accounts, bills, income sources, mortgages, HELOCs,
bill groups and promotional offers.

File format
-----------
Every data sheet has a header row (row 1) and a description/placeholder row
(row 2); data starts on row 3.  Columns are matched by *header name*, so
columns may be reordered and older files (which lack newer columns) still
import.  The column layout of every sheet is defined once in ``SHEET_COLUMNS``
and shared by the template generator, the exporter and the importer.

Template example rows carry a value in the ``is_example`` column (which the
exporter never writes); the importer skips any row flagged there.
"""

import inspect
import json
import os
from datetime import date, datetime, timedelta
from typing import Optional

# Try to import openpyxl, provide helpful message if not available
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# Column written only by the template generator, marking example rows.
EXAMPLE_COLUMN = "is_example"
EXAMPLE_MARKER = "EXAMPLE"

# (header, required, description) per sheet, in column order.
SHEET_COLUMNS: dict[str, list[tuple[str, bool, str]]] = {
    "Banks": [
        ("bank_name", True, "Name of the bank/institution"),
        ("bank_type", False, "traditional, credit_issuer, or mortgage_lender"),
    ],
    "Bank_Accounts": [
        ("bank_name", True, "Must match a name in Banks sheet"),
        ("account_name", True, "Name for this account"),
        ("account_type", True, "checking or savings"),
        ("balance", True, "Current balance"),
        ("interest_rate", False, "APR % for savings accounts"),
        ("interest_accrual_date", False, "Day of month (1-31) for interest"),
    ],
    "Credit_Accounts": [
        ("bank_name", True, "Must match a name in Banks sheet"),
        ("account_name", True, "Name for this credit account"),
        ("credit_limit", True, "Credit limit"),
        ("balance", True, "Current total balance owed (regular + promo)"),
        ("interest_rate", True, "APR %"),
        ("payment_due_date", False, "Day of month (1-31)"),
        ("statement_date", False, "Day of month (1-31)"),
        ("status", False, "active, closed, or frozen (default active)"),
    ],
    "Mortgages": [
        ("bank_name", True, "Lender name (will be created if not in Banks)"),
        ("account_name", True, "Name for this mortgage"),
        ("original_principal", True, "Original loan amount"),
        ("current_principal", True, "Current principal balance"),
        ("interest_rate", True, "APR %"),
        ("loan_term_months", True, "Total term (e.g., 360 for 30-year)"),
        ("payment_due_date", False, "Day of month (1-31)"),
        ("property_tax_monthly", False, "Monthly escrow for property tax"),
        ("insurance_monthly", False, "Monthly escrow for insurance"),
        ("pmi_monthly", False, "Monthly PMI if applicable"),
        ("hoa_monthly", False, "Monthly HOA fees if applicable"),
        ("other_monthly", False, "Other monthly escrow amounts"),
        ("start_date", False, "Loan start date (YYYY-MM-DD)"),
        ("escrow_balance", False, "Current escrow account balance"),
        ("market_value", False, "Current market value of the property"),
        ("preferred_source_bank", False, "Bank to pay the mortgage from"),
        ("preferred_source_account", False, "Account to pay the mortgage from"),
        ("status", False, "active, paid_off, refinanced, or sold"),
    ],
    "HELOCs": [
        ("bank_name", True, "Lender name (will be created if not in Banks)"),
        ("account_name", True, "Name for this HELOC"),
        ("credit_limit", True, "Credit limit"),
        ("balance", True, "Current amount drawn / owed"),
        ("interest_rate", True, "APR %"),
        ("draw_period_months", True, "Draw period length (e.g., 120)"),
        ("total_term_months", True, "Total term, longer than draw period (e.g., 240)"),
        ("start_date", False, "Date the HELOC was opened (YYYY-MM-DD)"),
        ("payment_due_date", False, "Day of month (1-31)"),
        ("preferred_source_bank", False, "Bank to pay the HELOC from"),
        ("preferred_source_account", False, "Account to pay the HELOC from"),
    ],
    "Bills": [
        ("bill_name", True, "Name of the bill"),
        ("amount", True, "Bill amount (estimated amount for floating bills)"),
        ("due_date", True, "Day of month (1-31)"),
        ("due_month", False, "Month (1-12) for yearly/semi-annual bills"),
        ("frequency", False, "once, weekly, biweekly, monthly, quarterly, semi_annual, yearly"),
        ("is_floating", False, "TRUE if amount varies, FALSE if fixed"),
        ("min_amount", False, "Minimum amount (for floating bills)"),
        ("max_amount", False, "Maximum amount (for floating bills)"),
        ("preferred_bank", False, "Preferred bank to pay from"),
        ("preferred_account", False, "Preferred account to pay from"),
        ("start_date", False, "First due date for non-monthly bills (YYYY-MM-DD)"),
        ("actual_amount", False, "Floating bills: actual amount of the current bill, if known"),
        ("amount_history", False, "Floating bills: past actual amounts, comma separated"),
    ],
    "Income": [
        ("source_name", True, "Name of income source"),
        ("amount", True, "Amount per pay period"),
        ("frequency", True, "once, weekly, biweekly, monthly, semimonthly, bimonthly, quarterly, yearly"),
        ("pay_date_1", True, "First pay date (day 1-31)"),
        ("pay_date_2", False, "Second pay date for semimonthly"),
        ("deposit_account", False, "Account to deposit into (ignored if distribution is set)"),
        ("deposit_bank", False, "Bank for deposit account (informational)"),
        ("distribution", False, 'JSON, e.g. {"Checking": {"value": 0.8, "type": "percentage"}}'),
        ("start_date", False, "First pay date for non-monthly income (YYYY-MM-DD)"),
    ],
    "Promo_Offers": [
        ("credit_account_name", True, "Must match a Credit Account name"),
        ("promo_name", True, "Name for this promotion"),
        ("months_interest_free", True, "Number of months at 0%"),
        ("purchase_limit", False, "Maximum amount for promo"),
        ("current_balance", False, "Current balance on this promo"),
        ("start_date", False, "When promo started (YYYY-MM-DD)"),
        ("end_date", False, "When promo ends (YYYY-MM-DD; computed if blank)"),
        ("amount_used", False, "Amount charged to the promo (defaults to current_balance)"),
        ("interest_charged", False, "TRUE if deferred interest was already charged"),
        ("bank_name", False, "Bank of the credit account (if the name is ambiguous)"),
    ],
    "Bill_Groups": [
        ("group_name", True, "Name of the bill group (one row per member)"),
        ("total_payment", True, "Total lump-sum payment"),
        ("due_date", True, "Day of month (1-31)"),
        ("frequency", False, "monthly, quarterly, ... (default monthly)"),
        ("group_type", False, "student_loan, bundled_bill, ..."),
        ("preferred_bank", False, "Bank to pay the group from"),
        ("preferred_account", False, "Account to pay the group from"),
        ("member_bank", False, "Bank of a member credit account"),
        ("member_account", False, "Member credit account name"),
        ("allocation_method", False, "proportional, fixed, or equal"),
        ("fixed_amount", False, "Amount for this member when allocation_method is fixed"),
    ],
}

SHEET_DESCRIPTIONS = {
    "Banks": "Define your financial institutions",
    "Bank_Accounts": "Checking and savings accounts",
    "Credit_Accounts": "Credit cards and loans",
    "Mortgages": "Home loans with escrow details",
    "HELOCs": "Home equity lines of credit",
    "Bills": "Recurring bills and expenses",
    "Income": "Income sources and pay schedules",
    "Promo_Offers": "Promotional 0% APR offers on credit cards",
    "Bill_Groups": "Bundled payments split across credit accounts (e.g. student loans)",
}

# Example rows for the template: {sheet: [ {header: value} ]}
TEMPLATE_EXAMPLES: dict[str, list[dict]] = {
    "Banks": [
        {"bank_name": "Chase Bank", "bank_type": "traditional"},
        {"bank_name": "Discover", "bank_type": "credit_issuer"},
    ],
    "Bank_Accounts": [
        {"bank_name": "Chase Bank", "account_name": "Primary Checking",
         "account_type": "checking", "balance": 5000.00},
        {"bank_name": "Chase Bank", "account_name": "Savings", "account_type": "savings",
         "balance": 10000.00, "interest_rate": 4.5, "interest_accrual_date": 1},
    ],
    "Credit_Accounts": [
        {"bank_name": "Discover", "account_name": "Discover It", "credit_limit": 10000,
         "balance": 2500.50, "interest_rate": 24.99, "payment_due_date": 20,
         "statement_date": 15, "status": "active"},
    ],
    "Mortgages": [
        {"bank_name": "Wells Fargo Mortgage", "account_name": "Primary Home",
         "original_principal": 350000, "current_principal": 320000, "interest_rate": 6.5,
         "loan_term_months": 360, "payment_due_date": 1, "property_tax_monthly": 350,
         "insurance_monthly": 150, "pmi_monthly": 0, "hoa_monthly": 0,
         "start_date": "2020-05-01"},
    ],
    "HELOCs": [
        {"bank_name": "Wells Fargo Mortgage", "account_name": "Home Equity Line",
         "credit_limit": 50000, "balance": 12000, "interest_rate": 8.0,
         "draw_period_months": 120, "total_term_months": 240,
         "start_date": "2024-02-01", "payment_due_date": 15},
    ],
    "Bills": [
        {"bill_name": "Rent/Mortgage", "amount": 2000, "due_date": 1, "frequency": "monthly",
         "is_floating": False},
        {"bill_name": "Electric", "amount": 150, "due_date": 15, "frequency": "monthly",
         "is_floating": True, "min_amount": 80, "max_amount": 250},
        {"bill_name": "Car Insurance", "amount": 180, "due_date": 20, "frequency": "monthly",
         "is_floating": False},
        {"bill_name": "Property Tax", "amount": 1200, "due_date": 15, "due_month": 4,
         "frequency": "yearly", "is_floating": False},
        {"bill_name": "Netflix", "amount": 15.99, "due_date": 5, "frequency": "monthly",
         "is_floating": False},
    ],
    "Income": [
        {"source_name": "Salary", "amount": 3500, "frequency": "semimonthly",
         "pay_date_1": 15, "pay_date_2": 31, "deposit_account": "Primary Checking",
         "deposit_bank": "Chase Bank"},
        {"source_name": "Side Gig", "amount": 500, "frequency": "monthly", "pay_date_1": 1},
    ],
    "Promo_Offers": [
        {"credit_account_name": "Discover It", "promo_name": "0% APR 18 months",
         "months_interest_free": 18, "purchase_limit": 5000, "current_balance": 2000,
         "start_date": "2025-06-01"},
    ],
    "Bill_Groups": [
        {"group_name": "Student Loans", "total_payment": 300, "due_date": 10,
         "frequency": "monthly", "group_type": "student_loan",
         "preferred_bank": "Chase Bank", "preferred_account": "Primary Checking",
         "member_bank": "Discover", "member_account": "Discover It",
         "allocation_method": "proportional"},
    ],
}

# Order in which sheets are written (and imported, respecting dependencies).
SHEET_ORDER = ["Banks", "Bank_Accounts", "Credit_Accounts", "Mortgages", "HELOCs",
               "Bills", "Income", "Promo_Offers", "Bill_Groups"]


# ---------------------------------------------------------------------------
# Value parsing helpers
# ---------------------------------------------------------------------------

def _to_str(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_float(value, default=None) -> Optional[float]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    return float(value)


def _to_int(value, default=None) -> Optional[int]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    return int(float(value))


def _to_bool(value) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1
    if isinstance(value, str):
        return value.strip().upper() in ("TRUE", "YES", "1", "Y", "X", "FLOATING", EXAMPLE_MARKER)
    return False


def _to_date(value) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    return date.fromisoformat(text[:10])


def _date_str(value) -> Optional[str]:
    return value.isoformat() if isinstance(value, date) else None


def _enum_value(value):
    return getattr(value, "value", value)


def _accepts_kwarg(cls, name: str) -> bool:
    """Whether ``cls``'s constructor accepts keyword ``name``."""
    try:
        return name in inspect.signature(cls.__init__).parameters
    except (TypeError, ValueError):
        return False


class DataImportExport:
    """Handles bulk import/export of financial data via Excel files."""

    def __init__(self, manager=None):
        """
        Args:
            manager: FinanceManager instance (optional for template generation)
        """
        self.manager = manager

    @staticmethod
    def check_openpyxl():
        """Check if openpyxl is available."""
        if not OPENPYXL_AVAILABLE:
            print("\n⚠️  Excel import/export requires 'openpyxl' package.")
            print("Install it with: pip install openpyxl")
            print("   or: uv add openpyxl")
            return False
        return True

    # ======================================================================
    # Template
    # ======================================================================

    def generate_template(self, filepath: str = None) -> str:
        """
        Generate an Excel template file with all required sheets and headers.

        Args:
            filepath: Path to save the template (default: Desktop)

        Returns:
            Path to the generated template file
        """
        if not self.check_openpyxl():
            return None

        if filepath is None:
            filepath = os.path.join(
                os.path.expanduser("~"),
                "Desktop",
                f"budget_template_{date.today().strftime('%Y%m%d')}.xlsx"
            )

        wb = Workbook()

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        required_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        optional_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        example_font = Font(italic=True, color="808080")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # ========== INSTRUCTIONS SHEET ==========
        ws_instructions = wb.active
        ws_instructions.title = "Instructions"

        instructions = [
            "BUDGET DATA IMPORT TEMPLATE",
            "",
            "This template allows you to bulk import your financial data.",
            "Fill out each sheet with your information and save the file.",
            "Then use the CLI option 'Import from Excel' to load the data.",
            "",
            "SHEET DESCRIPTIONS:",
            "",
        ]
        instructions += [f"{i}. {name} - {SHEET_DESCRIPTIONS[name]}"
                         for i, name in enumerate(SHEET_ORDER, 1)]
        instructions += [
            "",
            "TIPS:",
            "• Orange columns are REQUIRED",
            "• Green columns are OPTIONAL",
            "• Columns are matched by header name; do not rename the headers",
            "• Dates should be in YYYY-MM-DD format or Excel date format",
            "• Amounts should be numbers without $ or commas",
            f"• Example rows are marked '{EXAMPLE_MARKER}' in the '{EXAMPLE_COLUMN}' column and are",
            "  skipped on import. Delete them, or clear that cell to import the row.",
            "• Existing items with the same name are skipped (not duplicated) on import",
        ]

        for row_num, text in enumerate(instructions, 1):
            cell = ws_instructions.cell(row=row_num, column=1, value=text)
            if row_num == 1:
                cell.font = Font(bold=True, size=16)
            elif text.endswith(":"):
                cell.font = Font(bold=True)

        ws_instructions.column_dimensions['A'].width = 90

        # ========== DATA SHEETS ==========
        for sheet_name in SHEET_ORDER:
            ws = wb.create_sheet(sheet_name)
            headers = SHEET_COLUMNS[sheet_name] + [
                (EXAMPLE_COLUMN, False, "Rows marked EXAMPLE are skipped on import")
            ]
            self._setup_sheet(ws, headers, header_font, header_fill,
                              required_fill, optional_fill, thin_border)
            col_index = {h[0]: i for i, h in enumerate(headers, 1)}
            for row_num, example in enumerate(TEMPLATE_EXAMPLES.get(sheet_name, []), 3):
                row = dict(example)
                row[EXAMPLE_COLUMN] = EXAMPLE_MARKER
                for header, value in row.items():
                    if value is not None:
                        ws.cell(row=row_num, column=col_index[header], value=value).font = example_font

        wb.save(filepath)
        return filepath

    def _setup_sheet(self, ws, headers, header_font, header_fill,
                     required_fill, optional_fill, thin_border):
        """Set up a worksheet with headers and formatting."""
        for col_num, (header_name, is_required, description) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')

            # Add description row
            desc_cell = ws.cell(row=2, column=col_num, value=description)
            desc_cell.font = Font(italic=True, size=9)
            desc_cell.fill = required_fill if is_required else optional_fill
            desc_cell.border = thin_border
            desc_cell.alignment = Alignment(wrap_text=True)

            ws.column_dimensions[get_column_letter(col_num)].width = max(len(header_name) + 5, 15)

        # Freeze the header rows
        ws.freeze_panes = 'A3'

    # ======================================================================
    # Import
    # ======================================================================

    def import_from_excel(self, filepath: str) -> dict:
        """
        Import financial data from an Excel file.

        Items that already exist in the manager (same name, same parent) are
        skipped and reported in ``warnings`` and the ``*_skipped`` counters.

        Args:
            filepath: Path to the Excel file to import

        Returns:
            Dictionary with import results and any errors
        """
        if not self.check_openpyxl():
            return {"success": False, "error": "openpyxl not installed"}

        if not os.path.exists(filepath):
            return {"success": False, "error": f"File not found: {filepath}"}

        results = {
            "success": True,
            "banks_created": 0,
            "bank_accounts_created": 0,
            "credit_accounts_created": 0,
            "mortgages_created": 0,
            "helocs_created": 0,
            "bills_created": 0,
            "floating_bills_created": 0,
            "income_sources_created": 0,
            "promos_created": 0,
            "bill_groups_created": 0,
            "banks_skipped": 0,
            "bank_accounts_skipped": 0,
            "credit_accounts_skipped": 0,
            "mortgages_skipped": 0,
            "helocs_skipped": 0,
            "bills_skipped": 0,
            "income_sources_skipped": 0,
            "promos_skipped": 0,
            "bill_groups_skipped": 0,
            "example_rows_skipped": 0,
            "errors": [],
            "warnings": []
        }

        try:
            wb = load_workbook(filepath)
        except Exception as e:
            results["success"] = False
            results["errors"].append(f"Failed to read Excel file: {str(e)}")
            return results

        try:
            # Credit accounts whose regular balance must be recomputed once all
            # of their promos are known.
            self._touched_credits = {}

            # Import in order of dependencies
            self._import_banks(wb, results)
            self._import_bank_accounts(wb, results)
            self._import_credit_accounts(wb, results)
            self._import_mortgages(wb, results)
            self._import_helocs(wb, results)
            self._import_promos(wb, results)
            self._finalize_credit_balances()
            self._import_bill_groups(wb, results)
            self._import_bills(wb, results)
            self._import_income(wb, results)

            # Rebuild auto-generated linked bills (accounts / groups) and simulators
            if hasattr(self.manager, "refresh"):
                try:
                    self.manager.refresh()
                except Exception as e:  # pragma: no cover - defensive
                    results["warnings"].append(f"Could not refresh linked bills: {e}")
        except Exception as e:
            results["errors"].append(f"Import failed: {str(e)}")

        if results["errors"]:
            results["success"] = False

        return results

    def import_data(self, filepath: str) -> dict:
        """Alias for import_from_excel."""
        return self.import_from_excel(filepath)

    # ---------------------------------------------------------------- helpers

    def _sheet_rows(self, wb, sheet_name: str, results: dict, required: bool = True):
        """
        Yield (row_num, values) for each data row of a sheet, where ``values``
        maps lower-cased header names to cell values.  Blank rows and template
        example rows are skipped.
        """
        if sheet_name not in wb.sheetnames:
            if required:
                results["warnings"].append(f"{sheet_name} sheet not found, skipping")
            return

        ws = wb[sheet_name]
        headers: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            name = _to_str(ws.cell(row=1, column=col).value)
            if name:
                headers.setdefault(name.lower(), col)

        for row_num in range(3, ws.max_row + 1):  # Row 1 header, row 2 description
            values = {}
            for name, col in headers.items():
                value = ws.cell(row=row_num, column=col).value
                if isinstance(value, str) and not value.strip():
                    value = None
                values[name] = value
            if all(v is None for v in values.values()):
                continue
            if _to_bool(values.get(EXAMPLE_COLUMN)):
                results["example_rows_skipped"] += 1
                continue
            yield row_num, values

    def _find_bank(self, name):
        return next((b for b in self.manager.banks if b.name == name), None)

    def _find_or_create_bank(self, name, bank_type, results):
        from models.bank import Bank
        bank = self._find_bank(name)
        if not bank:
            bank = Bank(name, bank_type)
            self.manager.banks.append(bank)
            results["banks_created"] += 1
        return bank

    def _skip(self, results, kind: str, message: str):
        results[f"{kind}_skipped"] += 1
        results["warnings"].append(message)

    def _reserved_auto_bill_names(self) -> set[str]:
        """Names of bills that are auto-generated from accounts / bill groups."""
        names = set()
        for bank in self.manager.banks:
            for collection in (bank.credit_accounts, bank.mortgage_accounts,
                               getattr(bank, "heloc_accounts", {})):
                for acc_name in collection:
                    names.add(f"[Auto] {acc_name} Payment")
        for group in getattr(self.manager, "bill_groups", []):
            names.add(f"[Auto] {group.group_name} Payment")
        return names

    # ----------------------------------------------------------------- sheets

    def _import_banks(self, wb, results: dict):
        """Import banks from the Banks sheet."""
        from models.bank import Bank

        for row_num, row in self._sheet_rows(wb, "Banks", results):
            bank_name = _to_str(row.get("bank_name"))
            bank_type = _to_str(row.get("bank_type")) or "traditional"
            if not bank_name:
                continue
            try:
                if self._find_bank(bank_name):
                    self._skip(results, "banks", f"Bank '{bank_name}' already exists, skipping")
                    continue
                self.manager.banks.append(Bank(bank_name, bank_type))
                results["banks_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create bank '{bank_name}': {str(e)}")

    def _import_bank_accounts(self, wb, results: dict):
        """Import bank accounts from the Bank_Accounts sheet."""
        for row_num, row in self._sheet_rows(wb, "Bank_Accounts", results):
            bank_name = _to_str(row.get("bank_name"))
            account_name = _to_str(row.get("account_name"))
            if not bank_name or not account_name:
                continue
            try:
                bank = self._find_bank(bank_name)
                if not bank:
                    results["errors"].append(f"Row {row_num}: Bank '{bank_name}' not found")
                    continue
                if account_name in bank.accounts:
                    self._skip(results, "bank_accounts",
                               f"Account '{account_name}' already exists in '{bank_name}', skipping")
                    continue
                bank.add_account(
                    account_type=_to_str(row.get("account_type")) or "checking",
                    account_name=account_name,
                    starting_balance=_to_float(row.get("balance"), 0.0),
                    interest_rate=_to_float(row.get("interest_rate"), 0.0),
                    interest_accrual_date=_to_int(row.get("interest_accrual_date")),
                )
                results["bank_accounts_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create account '{account_name}': {str(e)}")

    def _import_credit_accounts(self, wb, results: dict):
        """Import credit accounts from the Credit_Accounts sheet."""
        from models.types import CreditAccountStatus

        for row_num, row in self._sheet_rows(wb, "Credit_Accounts", results):
            bank_name = _to_str(row.get("bank_name"))
            account_name = _to_str(row.get("account_name"))
            if not bank_name or not account_name:
                continue
            try:
                bank = self._find_bank(bank_name)
                if not bank:
                    results["errors"].append(f"Row {row_num}: Bank '{bank_name}' not found")
                    continue
                if account_name in bank.credit_accounts:
                    self._skip(results, "credit_accounts",
                               f"Credit account '{account_name}' already exists in '{bank_name}', skipping")
                    continue

                credit = bank.add_credit_account(
                    account_name=account_name,
                    credit_limit=_to_float(row.get("credit_limit"), 0.0),
                    interest_rate=_to_float(row.get("interest_rate"), 0.0),
                    payment_due_date=_to_int(row.get("payment_due_date")),
                    statement_date=_to_int(row.get("statement_date")),
                )

                balance = _to_float(row.get("balance"), 0.0)
                credit.balance = balance
                credit.regular_balance = balance

                status = _to_str(row.get("status"))
                if status:
                    credit.status = CreditAccountStatus(status.lower())

                self._touched_credits[id(credit)] = credit
                results["credit_accounts_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create credit account '{account_name}': {str(e)}")

    def _import_mortgages(self, wb, results: dict):
        """Import mortgages from the Mortgages sheet."""
        from models.mortgage_account import EscrowBreakdown
        from models.types import MortgageStatus

        for row_num, row in self._sheet_rows(wb, "Mortgages", results):
            bank_name = _to_str(row.get("bank_name"))
            account_name = _to_str(row.get("account_name"))
            if not bank_name or not account_name:
                continue
            try:
                bank = self._find_or_create_bank(bank_name, "mortgage_lender", results)
                if account_name in bank.mortgage_accounts:
                    self._skip(results, "mortgages",
                               f"Mortgage '{account_name}' already exists in '{bank_name}', skipping")
                    continue

                escrow = EscrowBreakdown(
                    property_tax_monthly=_to_float(row.get("property_tax_monthly"), 0.0),
                    homeowners_insurance_monthly=_to_float(row.get("insurance_monthly"), 0.0),
                    pmi_monthly=_to_float(row.get("pmi_monthly"), 0.0),
                    hoa_monthly=_to_float(row.get("hoa_monthly"), 0.0),
                    other_monthly=_to_float(row.get("other_monthly"), 0.0),
                )

                mortgage = bank.add_mortgage_account(
                    account_name=account_name,
                    original_principal=_to_float(row.get("original_principal"), 0.0),
                    current_principal=_to_float(row.get("current_principal"), 0.0),
                    interest_rate=_to_float(row.get("interest_rate"), 0.0),
                    loan_term_months=_to_int(row.get("loan_term_months"), 360),
                    payment_due_date=_to_int(row.get("payment_due_date"), 1),
                    escrow=escrow,
                    preferred_source_bank=_to_str(row.get("preferred_source_bank")),
                    preferred_source_account=_to_str(row.get("preferred_source_account")),
                )

                start_date = _to_date(row.get("start_date"))
                if start_date:
                    mortgage.start_date = start_date
                escrow_balance = _to_float(row.get("escrow_balance"))
                if escrow_balance is not None:
                    mortgage.escrow_balance = escrow_balance
                market_value = _to_float(row.get("market_value"))
                if market_value is not None:
                    mortgage.market_value = market_value
                status = _to_str(row.get("status"))
                if status:
                    mortgage.status = MortgageStatus(status.lower())

                results["mortgages_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create mortgage '{account_name}': {str(e)}")

    def _import_helocs(self, wb, results: dict):
        """Import HELOC accounts from the HELOCs sheet (optional sheet)."""
        for row_num, row in self._sheet_rows(wb, "HELOCs", results, required=False):
            bank_name = _to_str(row.get("bank_name"))
            account_name = _to_str(row.get("account_name"))
            if not bank_name or not account_name:
                continue
            try:
                bank = self._find_or_create_bank(bank_name, "mortgage_lender", results)
                if account_name in bank.heloc_accounts:
                    self._skip(results, "helocs",
                               f"HELOC '{account_name}' already exists in '{bank_name}', skipping")
                    continue
                bank.add_heloc_account(
                    account_name=account_name,
                    credit_limit=_to_float(row.get("credit_limit"), 0.0),
                    interest_rate=_to_float(row.get("interest_rate"), 0.0),
                    draw_period_months=_to_int(row.get("draw_period_months"), 180),
                    total_term_months=_to_int(row.get("total_term_months"), 360),
                    start_date=_to_date(row.get("start_date")),
                    current_balance=_to_float(row.get("balance"), 0.0),
                    payment_due_date=_to_int(row.get("payment_due_date"), 1),
                    preferred_source_bank=_to_str(row.get("preferred_source_bank")),
                    preferred_source_account=_to_str(row.get("preferred_source_account")),
                )
                results["helocs_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create HELOC '{account_name}': {str(e)}")

    def _import_promos(self, wb, results: dict):
        """Import promotional offers from the Promo_Offers sheet."""
        for row_num, row in self._sheet_rows(wb, "Promo_Offers", results):
            credit_name = _to_str(row.get("credit_account_name"))
            promo_name = _to_str(row.get("promo_name"))
            if not credit_name or not promo_name:
                continue
            try:
                bank_name = _to_str(row.get("bank_name"))
                credit = None
                for bank in self.manager.banks:
                    if bank_name and bank.name != bank_name:
                        continue
                    credit = bank.get_credit_account(credit_name)
                    if credit:
                        break
                if not credit:
                    results["errors"].append(f"Row {row_num}: Credit account '{credit_name}' not found")
                    continue
                if any(p.promo_name == promo_name for p in credit.active_promos):
                    self._skip(results, "promos",
                               f"Promo '{promo_name}' already exists on '{credit_name}', skipping")
                    continue

                months = _to_int(row.get("months_interest_free"), 12)
                start_date = _to_date(row.get("start_date"))
                end_date = _to_date(row.get("end_date"))
                if start_date is None and end_date is not None:
                    # Older exports wrote only end_date: derive the start.
                    start_date = end_date - timedelta(days=30 * months)

                promo = credit.add_promotional_offer(
                    promo_name=promo_name,
                    months_interest_free=months,
                    purchase_limit=_to_float(row.get("purchase_limit")),
                    start_date=start_date,
                )
                if end_date is not None:
                    promo.end_date = end_date

                balance = _to_float(row.get("current_balance"), 0.0)
                credit.promo_balance[promo_name] = balance
                amount_used = _to_float(row.get("amount_used"))
                promo.amount_used = amount_used if amount_used is not None else balance
                promo.interest_charged = _to_bool(row.get("interest_charged"))

                self._touched_credits[id(credit)] = credit
                results["promos_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create promo '{promo_name}': {str(e)}")

    def _finalize_credit_balances(self):
        """Regular balance = total balance minus the sum of all promo balances."""
        for credit in self._touched_credits.values():
            promo_total = sum(credit.promo_balance.values())
            credit.regular_balance = max(0.0, credit.balance - promo_total)
        self._touched_credits = {}

    def _import_bill_groups(self, wb, results: dict):
        """Import bill groups (one row per member) from the Bill_Groups sheet (optional)."""
        from models.bill_group import BillGroup, BillGroupMember
        from models.types import AllocationMethod, BillFrequency

        if not hasattr(self.manager, "bill_groups"):
            return

        new_groups: dict[str, BillGroup] = {}
        skipped: set[str] = set()
        for row_num, row in self._sheet_rows(wb, "Bill_Groups", results, required=False):
            group_name = _to_str(row.get("group_name"))
            if not group_name or group_name in skipped:
                continue
            try:
                group = new_groups.get(group_name)
                if group is None:
                    if any(g.group_name == group_name for g in self.manager.bill_groups):
                        skipped.add(group_name)
                        self._skip(results, "bill_groups",
                                   f"Bill group '{group_name}' already exists, skipping")
                        continue
                    frequency = _to_str(row.get("frequency")) or "monthly"
                    group = BillGroup(
                        group_name=group_name,
                        total_payment=_to_float(row.get("total_payment"), 0.0),
                        due_date=_to_int(row.get("due_date"), 1),
                        frequency=BillFrequency(frequency.lower()),
                        group_type=_to_str(row.get("group_type")) or "student_loan",
                        preferred_source_bank=_to_str(row.get("preferred_bank")),
                        preferred_source_account=_to_str(row.get("preferred_account")),
                    )
                    new_groups[group_name] = group

                member_account = _to_str(row.get("member_account"))
                if member_account:
                    member_bank = _to_str(row.get("member_bank"))
                    bank = self._find_bank(member_bank) if member_bank else next(
                        (b for b in self.manager.banks if member_account in b.credit_accounts), None)
                    if not bank or member_account not in bank.credit_accounts:
                        results["errors"].append(
                            f"Row {row_num}: Group member credit account '{member_account}' not found")
                        continue
                    method = _to_str(row.get("allocation_method")) or "proportional"
                    group.members.append(BillGroupMember(
                        bank_name=bank.name,
                        account_name=member_account,
                        allocation_method=AllocationMethod(method.lower()),
                        fixed_amount=_to_float(row.get("fixed_amount"), 0.0),
                    ))
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to import bill group '{group_name}': {str(e)}")

        for group in new_groups.values():
            self.manager.bill_groups.append(group)
            results["bill_groups_created"] += 1

    def _import_bills(self, wb, results: dict):
        """Import bills from the Bills sheet."""
        from models.fixed_bill import Bill, FloatingBill

        existing_names = {b.bill_name for b in self.manager.bills}
        auto_names = self._reserved_auto_bill_names()

        for row_num, row in self._sheet_rows(wb, "Bills", results):
            bill_name = _to_str(row.get("bill_name"))
            amount = _to_float(row.get("amount"))
            if not bill_name or not amount:
                continue
            try:
                if bill_name in auto_names:
                    self._skip(results, "bills",
                               f"Bill '{bill_name}' is auto-generated from an account or bill group, skipping")
                    continue
                if bill_name in existing_names:
                    self._skip(results, "bills", f"Bill '{bill_name}' already exists, skipping")
                    continue

                frequency = _to_str(row.get("frequency")) or "monthly"
                min_amount = _to_float(row.get("min_amount"))
                max_amount = _to_float(row.get("max_amount"))
                # Treat as floating if flagged, or if min/max amounts are specified
                is_floating = _to_bool(row.get("is_floating")) or \
                    min_amount is not None or max_amount is not None

                kwargs = dict(
                    bill_name=bill_name,
                    due_date=_to_int(row.get("due_date"), 1),
                    frequency=frequency,
                    due_month=_to_int(row.get("due_month")),
                    preferred_source_bank=_to_str(row.get("preferred_bank")),
                    preferred_source_account=_to_str(row.get("preferred_account")),
                )
                cls = FloatingBill if is_floating else Bill
                start_date = _to_date(row.get("start_date"))
                if start_date is not None and _accepts_kwarg(cls, "start_date"):
                    kwargs["start_date"] = start_date

                if is_floating:
                    bill = FloatingBill(
                        estimated_amount=amount,
                        min_amount=min_amount if min_amount is not None else amount * 0.5,
                        max_amount=max_amount if max_amount is not None else amount * 1.5,
                        **kwargs,
                    )
                    history = row.get("amount_history")
                    if history not in (None, ""):
                        bill.amount_history = [float(x) for x in str(history).split(",") if x.strip()]
                    actual = _to_float(row.get("actual_amount"))
                    if actual is not None:
                        bill.set_actual_amount(actual)
                    results["floating_bills_created"] += 1
                else:
                    bill = Bill(amount=amount, **kwargs)

                self.manager.bills.append(bill)
                existing_names.add(bill_name)
                results["bills_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create bill '{bill_name}': {str(e)}")

    def _import_income(self, wb, results: dict):
        """Import income sources from the Income sheet."""
        from models.income_source import IncomeSource, IncomeFrequency

        existing_names = {i.source_name for i in self.manager.income_sources}

        for row_num, row in self._sheet_rows(wb, "Income", results):
            source_name = _to_str(row.get("source_name"))
            amount = _to_float(row.get("amount"))
            if not source_name or not amount:
                continue
            try:
                if source_name in existing_names:
                    self._skip(results, "income_sources",
                               f"Income source '{source_name}' already exists, skipping")
                    continue

                pay_dates = [d for d in (_to_int(row.get("pay_date_1")),
                                         _to_int(row.get("pay_date_2"))) if d]
                if not pay_dates:
                    pay_dates = [1]  # Default to 1st of month

                distribution = None
                dist_json = row.get("distribution")
                if dist_json not in (None, ""):
                    distribution = json.loads(str(dist_json)) if not isinstance(dist_json, dict) else dist_json
                else:
                    deposit_account = _to_str(row.get("deposit_account"))
                    if deposit_account:
                        distribution = {deposit_account: amount}

                frequency = _to_str(row.get("frequency")) or "monthly"
                kwargs = dict(
                    source_name=source_name,
                    amount=amount,
                    frequency=IncomeFrequency(frequency.lower()),
                    pay_date=pay_dates,
                    distribution=distribution,
                )
                start_date = _to_date(row.get("start_date"))
                if start_date is not None and _accepts_kwarg(IncomeSource, "start_date"):
                    kwargs["start_date"] = start_date

                self.manager.income_sources.append(IncomeSource(**kwargs))
                existing_names.add(source_name)
                results["income_sources_created"] += 1
            except Exception as e:
                results["errors"].append(f"Row {row_num}: Failed to create income '{source_name}': {str(e)}")

    # ======================================================================
    # Export
    # ======================================================================

    def export_current_data(self, filepath: str = None) -> str:
        """
        Export current financial data to an Excel file that
        ``import_from_excel`` can read back.

        Args:
            filepath: Path to save the export (default: Desktop)

        Returns:
            Path to the exported file
        """
        if not self.check_openpyxl():
            return None

        if filepath is None:
            filepath = os.path.join(
                os.path.expanduser("~"),
                "Desktop",
                f"budget_export_{date.today().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )

        rows = {name: [] for name in SHEET_ORDER}
        m = self.manager

        for bank in m.banks:
            rows["Banks"].append({"bank_name": bank.name, "bank_type": _enum_value(bank.bank_type)})

            for name, account in bank.accounts.items():
                rows["Bank_Accounts"].append({
                    "bank_name": bank.name,
                    "account_name": name,
                    "account_type": _enum_value(account.account_type),
                    "balance": account.balance,
                    "interest_rate": account.interest_rate,
                    "interest_accrual_date": getattr(account, "interest_accrual_date", None),
                })

            for name, credit in bank.credit_accounts.items():
                rows["Credit_Accounts"].append({
                    "bank_name": bank.name,
                    "account_name": name,
                    "credit_limit": credit.credit_limit,
                    "balance": credit.balance,
                    "interest_rate": credit.interest_rate,
                    "payment_due_date": credit.payment_due_date,
                    "statement_date": credit.statement_date,
                    "status": _enum_value(getattr(credit, "status", None)),
                })
                for promo in credit.active_promos:
                    rows["Promo_Offers"].append({
                        "credit_account_name": name,
                        "promo_name": promo.promo_name,
                        "months_interest_free": promo.months_interest_free,
                        "purchase_limit": promo.purchase_limit,
                        "current_balance": credit.promo_balance.get(promo.promo_name, 0),
                        "start_date": _date_str(promo.start_date),
                        "end_date": _date_str(promo.end_date),
                        "amount_used": promo.amount_used,
                        "interest_charged": bool(promo.interest_charged),
                        "bank_name": bank.name,
                    })

            for name, mortgage in bank.mortgage_accounts.items():
                escrow = mortgage.escrow
                rows["Mortgages"].append({
                    "bank_name": bank.name,
                    "account_name": name,
                    "original_principal": mortgage.original_principal,
                    "current_principal": mortgage.current_principal,
                    "interest_rate": mortgage.interest_rate,
                    "loan_term_months": mortgage.loan_term_months,
                    "payment_due_date": mortgage.payment_due_date,
                    "property_tax_monthly": escrow.property_tax_monthly,
                    "insurance_monthly": escrow.homeowners_insurance_monthly,
                    "pmi_monthly": escrow.pmi_monthly,
                    "hoa_monthly": escrow.hoa_monthly,
                    "other_monthly": escrow.other_monthly,
                    "start_date": _date_str(mortgage.start_date),
                    "escrow_balance": mortgage.escrow_balance,
                    "market_value": mortgage.market_value,
                    "preferred_source_bank": mortgage.preferred_source_bank,
                    "preferred_source_account": mortgage.preferred_source_account,
                    "status": _enum_value(mortgage.status),
                })

            for name, heloc in getattr(bank, "heloc_accounts", {}).items():
                rows["HELOCs"].append({
                    "bank_name": bank.name,
                    "account_name": name,
                    "credit_limit": heloc.credit_limit,
                    "balance": heloc.balance,
                    "interest_rate": heloc.interest_rate,
                    "draw_period_months": heloc.draw_period_months,
                    "total_term_months": heloc.total_term_months,
                    "start_date": _date_str(heloc.start_date),
                    "payment_due_date": heloc.payment_due_date,
                    "preferred_source_bank": heloc.preferred_source_bank,
                    "preferred_source_account": heloc.preferred_source_account,
                })

        for bill in m.bills:
            # Skip auto-generated bills — they are recreated from accounts / groups
            if getattr(bill, "is_account_linked", False) or getattr(bill, "is_group_linked", False):
                continue
            is_floating = hasattr(bill, "min_amount")
            row = {
                "bill_name": bill.bill_name,
                "amount": bill.estimated_amount if is_floating else bill.amount,
                "due_date": bill.due_date,
                "due_month": bill.due_month,
                "frequency": _enum_value(bill.frequency),
                "is_floating": is_floating,
                "preferred_bank": bill.preferred_source_bank,
                "preferred_account": bill.preferred_source_account,
                "start_date": _date_str(getattr(bill, "start_date", None)),
            }
            if is_floating:
                row["min_amount"] = bill.min_amount
                row["max_amount"] = bill.max_amount
                row["actual_amount"] = getattr(bill, "_actual_amount", None)
                history = getattr(bill, "amount_history", None) or []
                row["amount_history"] = ",".join(repr(float(x)) for x in history) or None
            rows["Bills"].append(row)

        for income in m.income_sources:
            pay_dates = income.pay_date or []
            dist = income.distribution
            dist_entries = list(dist.distribution) if dist else []
            rows["Income"].append({
                "source_name": income.source_name,
                "amount": income.amount,
                "frequency": _enum_value(income.frequency),
                "pay_date_1": pay_dates[0] if len(pay_dates) > 0 else None,
                "pay_date_2": pay_dates[1] if len(pay_dates) > 1 else None,
                # Human-readable hint; the distribution column is authoritative
                "deposit_account": dist_entries[0].account_name if len(dist_entries) == 1 else None,
                "distribution": json.dumps(dist.to_json()) if dist else None,
                "start_date": _date_str(getattr(income, "start_date", None)),
            })

        for group in getattr(m, "bill_groups", []):
            base = {
                "group_name": group.group_name,
                "total_payment": group.total_payment,
                "due_date": group.due_date,
                "frequency": _enum_value(group.frequency),
                "group_type": group.group_type,
                "preferred_bank": group.preferred_source_bank,
                "preferred_account": group.preferred_source_account,
            }
            if not group.members:
                rows["Bill_Groups"].append(base)
            for member in group.members:
                rows["Bill_Groups"].append({
                    **base,
                    "member_bank": member.bank_name,
                    "member_account": member.account_name,
                    "allocation_method": _enum_value(member.allocation_method),
                    "fixed_amount": member.fixed_amount,
                })

        wb = Workbook()
        wb.remove(wb.active)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        desc_font = Font(italic=True, size=9, color="808080")

        for sheet_name in SHEET_ORDER:
            ws = wb.create_sheet(sheet_name)
            headers = [h[0] for h in SHEET_COLUMNS[sheet_name]]
            ws.append(headers)
            # Row 2 placeholder (template description row) for import compatibility
            ws.append(["(data exported)"] + [""] * (len(headers) - 1))
            for col in range(1, len(headers) + 1):
                ws.cell(row=1, column=col).font = header_font
                ws.cell(row=1, column=col).fill = header_fill
                ws.column_dimensions[get_column_letter(col)].width = max(len(headers[col - 1]) + 5, 15)
            ws.cell(row=2, column=1).font = desc_font
            ws.freeze_panes = 'A3'
            for row in rows[sheet_name]:
                ws.append([row.get(h) for h in headers])

        wb.save(filepath)
        return filepath
